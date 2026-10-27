"""Excel export of Weekly Wage Sheets: a pay sheet per week, with a column per day."""

import re
from io import BytesIO

import frappe
from frappe import _
from frappe.utils import add_days, date_diff, flt, format_date, getdate
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BOLD = Font(bold=True)
TITLE = Font(bold=True, size=14)
HEAD_FILL = PatternFill("solid", fgColor="DDE4EE")
TOTAL_FILL = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MONEY = "#,##0.00"
DAYS = "0.##"
STATUS_LABEL = {0: "Draft", 1: "Submitted", 2: "Cancelled"}


@frappe.whitelist()
def export_wage_sheets(names):
	names = frappe.parse_json(names) if isinstance(names, str) and names.startswith("[") else names
	names = [names] if isinstance(names, str) else list(names)
	if not names:
		frappe.throw(_("Select at least one wage sheet"))

	sheets = []
	for name in names:
		doc = frappe.get_doc("Weekly Wage Sheet", name)
		doc.check_permission("read")
		sheets.append(doc)
	sheets.sort(key=lambda d: (d.project, getdate(d.from_date)))

	wb = Workbook()
	wb.remove(wb.active)
	if len(sheets) > 1:
		_write_summary(wb.create_sheet(_("Summary")), sheets)
	used = set()
	for doc in sheets:
		_write_pay_sheet(wb.create_sheet(_sheet_title(doc, used)), doc)

	out = BytesIO()
	wb.save(out)
	if len(sheets) == 1:
		d = sheets[0]
		filename = f"Wages {d.project} {d.from_date} to {d.to_date}.xlsx"
	else:
		filename = f"Wages {min(d.from_date for d in sheets)} to {max(d.to_date for d in sheets)}.xlsx"
	frappe.response.filename = re.sub(r'[\\/:*?"<>|]', "-", filename)
	frappe.response.filecontent = out.getvalue()
	frappe.response.type = "binary"


def _sheet_title(doc, used):
	base = re.sub(r"[\\/*?:\[\]]", "-", f"{doc.project} {format_date(doc.from_date, 'dd-MM')}")[:28]
	title, n = base, 2
	while title in used:
		title, n = f"{base[:25]} ({n})", n + 1
	used.add(title)
	return title


def _attendance(doc):
	"""Per worker, per date: (status, day_fraction, daily_rate, wage) from the registers this sheet covers."""
	registers = [n for n in (doc.attendance_list or "").split("\n") if n]
	if not registers:
		return {}
	rows = frappe.db.sql(
		"""
		select sa.attendance_date, e.worker, e.role, e.status, e.day_fraction, e.daily_rate, e.wage
		from `tabSite Attendance` sa
		join `tabSite Attendance Entry` e on e.parent = sa.name and e.parenttype = 'Site Attendance'
		where sa.name in %(registers)s
		""",
		{"registers": registers},
		as_dict=True,
	)
	out = {}
	for r in rows:
		out.setdefault(r.worker, {})[getdate(r.attendance_date)] = r
	return out


def _write_pay_sheet(ws, doc):
	att = _attendance(doc)
	dates = [getdate(add_days(doc.from_date, i)) for i in range(date_diff(doc.to_date, doc.from_date) + 1)]
	workers = frappe.get_all(
		"Site Worker",
		filters={"name": ["in", [e.worker for e in doc.entries] or [""]]},
		fields=["name", "id_number", "phone"],
	)
	info = {w.name: w for w in workers}
	company = doc.company or frappe.db.get_value("Project", doc.project, "company")

	ws["A1"] = company
	ws["A1"].font = TITLE
	ws["A2"] = _("Site Wages: {0} ({1})").format(doc.project_name or doc.project, doc.project)
	ws["A2"].font = BOLD
	ws["A3"] = _("Week {0} to {1}").format(format_date(doc.from_date), format_date(doc.to_date))
	ws["A4"] = _("Wage Sheet {0} · {1}{2}").format(
		doc.name,
		_(STATUS_LABEL.get(doc.docstatus, "")),
		_(" · Paid {0}").format(format_date(doc.payment_date)) if doc.docstatus == 1 and doc.payment_date else "",
	)

	headers = ["#", _("Worker Code"), _("Name"), _("ID Number"), _("M-Pesa Number"), _("Role")]
	headers += [format_date(d, "EEE dd-MM") for d in dates]
	headers += [_("Days"), _("Daily Rate"), _("Amount"), _("Payment Reference"), _("Signature")]
	first_day_col = 7
	days_col = first_day_col + len(dates)
	head_row = 6
	for c, h in enumerate(headers, start=1):
		cell = ws.cell(row=head_row, column=c, value=h)
		cell.font, cell.fill, cell.border = BOLD, HEAD_FILL, BOX
		cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

	r = head_row
	for i, e in enumerate(doc.entries, start=1):
		r += 1
		days = att.get(e.worker, {})
		roles = {a.role for a in days.values() if a.role}
		rates = {flt(a.daily_rate) for a in days.values() if flt(a.day_fraction)}
		w = info.get(e.worker) or frappe._dict()
		values = [i, e.worker, e.worker_name, w.id_number, e.mpesa_number or w.phone, ", ".join(sorted(roles))]
		for d in dates:
			a = days.get(d)
			values.append(flt(a.day_fraction) if a and flt(a.day_fraction) else None)
		rate = rates.pop() if len(rates) == 1 else (flt(e.amount) / flt(e.days_worked) if flt(e.days_worked) else None)
		values += [flt(e.days_worked), rate, flt(e.amount), e.payment_reference, None]
		for c, v in enumerate(values, start=1):
			cell = ws.cell(row=r, column=c, value=v)
			cell.border = BOX
		for c in range(first_day_col, days_col + 1):
			ws.cell(row=r, column=c).number_format = DAYS
			ws.cell(row=r, column=c).alignment = Alignment(horizontal="center")
		ws.cell(row=r, column=days_col + 1).number_format = MONEY
		ws.cell(row=r, column=days_col + 2).number_format = MONEY
		for c in (2, 4, 5, days_col + 3):  # keep codes and phone numbers as text
			ws.cell(row=r, column=c).number_format = "@"
		ws.row_dimensions[r].height = 22

	# totals (as formulas, so edits in Excel still add up)
	total_row = r + 1
	ws.cell(row=total_row, column=3, value=_("Total ({0} workers)").format(len(doc.entries)))
	cols = list(range(first_day_col, days_col + 1)) + [days_col + 2]
	for c in cols:
		letter = get_column_letter(c)
		value = f"=SUM({letter}{head_row + 1}:{letter}{r})" if r > head_row else 0
		ws.cell(row=total_row, column=c, value=value)
		ws.cell(row=total_row, column=c).number_format = MONEY if c == days_col + 2 else DAYS
	for c in range(1, len(headers) + 1):
		cell = ws.cell(row=total_row, column=c)
		cell.font, cell.fill, cell.border = BOLD, TOTAL_FILL, BOX

	ws.cell(row=total_row + 2, column=2, value=_("Day key: 1 = full day, 0.5 = half day, blank = not present."))
	ws.cell(row=total_row + 4, column=2, value=_("Prepared by: ____________________"))
	ws.cell(row=total_row + 4, column=days_col, value=_("Approved by: ____________________"))

	widths = [5, 12, 26, 14, 15, 14] + [9] * len(dates) + [8, 12, 13, 20, 18]
	for c, width in enumerate(widths, start=1):
		ws.column_dimensions[get_column_letter(c)].width = width
	ws.row_dimensions[head_row].height = 30
	ws.freeze_panes = ws.cell(row=head_row + 1, column=4)
	ws.page_setup.orientation = "landscape"
	ws.page_setup.fitToWidth = 1
	ws.page_setup.fitToHeight = 0
	ws.sheet_properties.pageSetUpPr.fitToPage = True
	ws.print_title_rows = f"{head_row}:{head_row}"


def _write_summary(ws, sheets):
	ws["A1"] = _("Site Wages Summary")
	ws["A1"].font = TITLE
	headers = [_("Wage Sheet"), _("Project"), _("Project Name"), _("Week From"), _("Week To"), _("Status"),
		_("Payment Date"), _("Workers"), _("Days"), _("Amount")]
	for c, h in enumerate(headers, start=1):
		cell = ws.cell(row=3, column=c, value=h)
		cell.font, cell.fill, cell.border = BOLD, HEAD_FILL, BOX
	r = 3
	for d in sheets:
		r += 1
		values = [d.name, d.project, d.project_name, getdate(d.from_date), getdate(d.to_date),
			_(STATUS_LABEL.get(d.docstatus, "")), getdate(d.payment_date) if d.docstatus == 1 and d.payment_date else None,
			d.total_workers, flt(d.total_days), flt(d.total_amount)]
		for c, v in enumerate(values, start=1):
			cell = ws.cell(row=r, column=c, value=v)
			cell.border = BOX
		for c in (4, 5, 7):
			ws.cell(row=r, column=c).number_format = "DD-MM-YYYY"
		ws.cell(row=r, column=9).number_format = DAYS
		ws.cell(row=r, column=10).number_format = MONEY
	ws.cell(row=r + 1, column=1, value=_("Total"))
	for c, fmt in ((8, "0"), (9, DAYS), (10, MONEY)):
		letter = get_column_letter(c)
		ws.cell(row=r + 1, column=c, value=f"=SUM({letter}4:{letter}{r})").number_format = fmt
	for c in range(1, len(headers) + 1):
		cell = ws.cell(row=r + 1, column=c)
		cell.font, cell.fill, cell.border = BOLD, TOTAL_FILL, BOX
	for c, width in enumerate([18, 14, 28, 12, 12, 11, 13, 9, 8, 14], start=1):
		ws.column_dimensions[get_column_letter(c)].width = width
	ws.freeze_panes = "A4"
